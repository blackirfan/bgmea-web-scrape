"""Excel output: a fresh workbook, or filling an existing template in place."""

import os
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from fields import BUSINESS_COLUMNS


class OutputLocked(RuntimeError):
    """The target .xlsx is open in another program (usually Excel)."""


def _save_atomic(wb, path):
    """Save via a temp file and swap it in.

    If `path` is locked (open in Excel), the data is written to
    `<name>.locked.xlsx` instead and OutputLocked is raised so the caller can
    warn without losing anything.
    """
    path = Path(path)
    tmp = path.with_suffix(path.suffix + ".tmp")
    wb.save(tmp)
    try:
        os.replace(tmp, path)
    except PermissionError:
        fallback = path.with_name(path.stem + ".locked" + path.suffix)
        try:
            os.replace(tmp, fallback)
        except PermissionError:
            fallback.unlink(missing_ok=True)
            os.replace(tmp, fallback)
        raise OutputLocked(
            f"{path.name} is open in another program; wrote {fallback.name} "
            f"instead. Close {path.name} to let the next checkpoint update it.")

_HEAD_FILL = PatternFill("solid", fgColor="1F3864")
_HEAD_FONT = Font(bold=True, color="FFFFFF")


def _style_header(ws):
    for cell in ws[1]:
        cell.font = _HEAD_FONT
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _autosize(ws, cap=48):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None),
                    default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(12, width + 2), cap)


def _reg_as_text(ws):
    idx = BUSINESS_COLUMNS.index("BGMEA registration no.") + 1
    for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
        for cell in row:
            cell.number_format = "@"
            if cell.value is not None:
                cell.value = str(cell.value)


def write_workbook(path, records):
    wb = Workbook()
    ws = wb.active
    ws.title = "Factories"
    ws.append(BUSINESS_COLUMNS)
    _style_header(ws)
    for rec in records:
        ws.append([rec.get(c) for c in BUSINESS_COLUMNS])
    _reg_as_text(ws)
    _autosize(ws)
    ws.freeze_panes = "A2"
    _save_atomic(wb, path)


def write_audit(path, audit_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit"
    ws.append(["Name of factory", "BGMEA registration no.", "BGMEA URL", "Notes"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in audit_rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["C"].width = 42
    ws.column_dimensions["D"].width = 90
    _save_atomic(wb, path)


# --------------------------------------------------------------------------- #
#  template filling                                                           #
# --------------------------------------------------------------------------- #

_LEGAL = {"limited": "ltd", "private": "pvt", "company": "co",
          "industries": "inds", "and": "&"}


def _name_key(name):
    text = re.sub(r"(?ix)\(?\s*(?:unit|u)\s*[-.\#]?\s*\d{1,3}\s*\)?", " ", name or "")
    text = re.sub(r"[^\w\s]", " ", text.lower())
    return " ".join(_LEGAL.get(t, t) for t in text.split())


def _reg_key(value):
    text = re.sub(r"\.0+$", "", str(value or "").strip())
    return text if re.fullmatch(r"\d+", text) else None


def write_template(template_path, output_path, records):
    """Copy `template_path`, fill columns whose header matches ours.

    Existing non-empty cells are never overwritten. Rows are matched by
    registration number first, then by a normalised factory name.
    """
    wb = load_workbook(template_path)
    ws = wb.active

    header_col = {}
    for cell in ws[1]:
        if cell.value:
            header_col[str(cell.value).strip().lower()] = cell.column
    targets = {c: header_col[c.lower()] for c in BUSINESS_COLUMNS
               if c.lower() in header_col}
    if not targets:
        write_workbook(output_path, records)
        return

    by_reg, by_name = {}, {}
    for rec in records:
        rk = _reg_key(rec.get("BGMEA registration no."))
        if rk:
            by_reg.setdefault(rk, rec)
        nk = _name_key(rec.get("Name of factory"))
        if nk:
            by_name.setdefault(nk, rec)

    name_c = header_col.get("name of factory")
    reg_c = header_col.get("bgmea registration no.")
    for r in range(2, ws.max_row + 1):
        src = None
        if reg_c:
            rk = _reg_key(ws.cell(r, reg_c).value)
            if rk:
                src = by_reg.get(rk)
        if src is None and name_c:
            src = by_name.get(_name_key(ws.cell(r, name_c).value))
        if src is None:
            continue
        for col_name, col_idx in targets.items():
            cell = ws.cell(r, col_idx)
            if cell.value not in (None, ""):
                continue
            value = src.get(col_name)
            if value in (None, ""):
                continue
            if col_name == "BGMEA registration no.":
                cell.number_format = "@"
                cell.value = str(value)
            else:
                cell.value = value
    _save_atomic(wb, output_path)
