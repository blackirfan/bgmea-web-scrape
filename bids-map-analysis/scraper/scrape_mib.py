"""
Scrapes all factory detail pages from mappedinbangladesh.org into a new
Excel file, using an already-open, logged-in Chrome you control yourself.

RUN IT MANUALLY, IN THE FOREGROUND:

    py -3 scraper/scrape_mib.py

Before running:
  1. Launch a dedicated Chrome window with remote debugging enabled:
       "C:\\Program Files\\Google\\Chrome\\Application\\chrome.exe" ^
         --remote-debugging-port=9222 --user-data-dir=C:\\Users\\hp\\chrome-mib-automation
  2. Log into mappedinbangladesh.org in THAT window (one-time only -- the
     profile folder keeps you logged in for future runs).
  3. Leave that window open, then run this script.

You can stop it any time with Ctrl+C and re-run later -- it resumes from
where it left off (see scraper/state/scraped.jsonl).
"""
import json
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

# Factory names/addresses contain non-ASCII characters (Bengali text, curly
# quotes, special hyphens); Windows' default console/file encoding (cp1252)
# can't print those and would crash the whole run. Force UTF-8 output.
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

import openpyxl
from selenium import webdriver
from selenium.common.exceptions import WebDriverException

from extract import OUTPUT_COLUMNS, extract_factory, is_placeholder_heavy

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent

DEBUG_ADDRESS = "127.0.0.1:9222"
SOURCE_XLSX = ROOT / "excel" / "mib_factories.xlsx"
OUTPUT_XLSX = ROOT / "excel" / "mib_factories_full.xlsx"
STATE_DIR = HERE / "state"
JSONL_PATH = STATE_DIR / "scraped.jsonl"

PAGE_LOAD_WAIT_TIMEOUT = 15
POLITE_DELAY_SECONDS = 1.0


def connect_driver():
    options = webdriver.ChromeOptions()
    options.debugger_address = DEBUG_ADDRESS
    try:
        return webdriver.Chrome(options=options)
    except WebDriverException as e:
        print("\nCould not connect to Chrome's debug port at", DEBUG_ADDRESS)
        print("Make sure the dedicated automation Chrome window is open (see the")
        print("instructions at the top of this script), then re-run.")
        print("Details:", e)
        sys.exit(1)


def load_source_rows():
    """Read (factory_id, factory_url, latitude, longitude) for every row."""
    wb = openpyxl.load_workbook(SOURCE_XLSX, read_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {name: i for i, name in enumerate(header)}
    out = []
    for r in rows[1:]:
        fid = r[idx["factory_id"]]
        url = r[idx["factory_url"]]
        lat = r[idx.get("latitude")] if "latitude" in idx else None
        lon = r[idx.get("longitude")] if "longitude" in idx else None
        if fid is None or url is None:
            continue
        out.append((str(fid), str(url), lat, lon))
    wb.close()
    return out


def load_done_ids():
    done = {}
    if not JSONL_PATH.exists():
        return done
    with open(JSONL_PATH, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except json.JSONDecodeError:
                continue
            done[rec["factory_id"]] = rec
    return done


def load_and_extract(driver, url, max_attempts=3):
    """
    Load a factory page and extract it, retrying if the page is still
    showing the "Not Reported" placeholder skeleton (the data loads via an
    async fetch after the initial page load, so a single quick read is not
    reliable).
    """
    record = None
    for attempt in range(1, max_attempts + 1):
        driver.get(url)
        deadline = time.time() + PAGE_LOAD_WAIT_TIMEOUT
        time.sleep(1.5)
        while True:
            record = extract_factory(driver)
            if record["scrape_status"] == "login_required":
                return record
            if not is_placeholder_heavy(record):
                return record
            if time.time() >= deadline:
                break
            time.sleep(1.0)
        # still placeholder-heavy after the full wait -- try a fresh load
    record["scrape_status"] = "placeholder_suspect"
    return record


def build_excel_from_jsonl():
    by_id = {}
    if JSONL_PATH.exists():
        with open(JSONL_PATH, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line:
                    rec = json.loads(line)
                    # Later attempts overwrite earlier ones for the same factory
                    # (e.g. a retried factory across resumed runs).
                    by_id[rec["factory_id"]] = rec
    records = list(by_id.values())
    records.sort(key=lambda r: int(r["factory_id"]) if r["factory_id"].isdigit() else 0)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(OUTPUT_COLUMNS)
    for rec in records:
        ws.append([rec.get(col, "") for col in OUTPUT_COLUMNS])
    wb.save(OUTPUT_XLSX)


def append_jsonl(record):
    STATE_DIR.mkdir(parents=True, exist_ok=True)
    with open(JSONL_PATH, "a", encoding="utf-8") as f:
        f.write(json.dumps(record, ensure_ascii=False))
        f.write("\n")


def main():
    limit = None
    if len(sys.argv) > 1 and sys.argv[1] == "--limit":
        limit = int(sys.argv[2])

    print("Reading factory URLs from", SOURCE_XLSX.relative_to(ROOT))
    rows = load_source_rows()
    total = len(rows)
    print(f"Found {total} factory URLs.")
    if limit:
        print(f"(--limit {limit}: this run will stop after {limit} new factories)")

    done = load_done_ids()
    already_ok = sum(1 for r in done.values() if r.get("scrape_status") == "ok")
    print(f"Already scraped (resumable): {len(done)} ({already_ok} ok, {len(done) - already_ok} need retry)")

    print(f"Connecting to Chrome at {DEBUG_ADDRESS} ...")
    driver = connect_driver()
    print("Connected. Starting scrape. Press Ctrl+C to stop at any time (safe to resume).\n")

    remaining_rows = [r for r in rows if r[0] not in done or done[r[0]].get("scrape_status") != "ok"]
    if limit:
        remaining_rows = remaining_rows[:limit]
    remaining = len(remaining_rows)
    scraped_this_run = 0
    login_pauses = 0

    try:
        for i, (fid, url, lat, lon) in enumerate(remaining_rows, start=1):
            t0 = time.time()
            try:
                record = load_and_extract(driver, url)
            except WebDriverException as e:
                record = {col: "" for col in OUTPUT_COLUMNS}
                record["scrape_status"] = f"webdriver_error: {e.__class__.__name__}"

            record["factory_id"] = fid
            record["factory_url"] = url
            record["latitude"] = lat if lat is not None else ""
            record["longitude"] = lon if lon is not None else ""
            record["scraped_at"] = datetime.now(timezone.utc).isoformat()

            while record["scrape_status"] == "login_required":
                login_pauses += 1
                print("\n" + "=" * 70)
                print("LOGIN REQUIRED: the site is showing a logged-out page.")
                print("Please check the automation Chrome window and log back in,")
                print("then press Enter here to retry this factory.")
                print("=" * 70)
                input(">> Press Enter once you're logged in again... ")
                record = load_and_extract(driver, url)
                record["factory_id"] = fid
                record["factory_url"] = url
                record["latitude"] = lat if lat is not None else ""
                record["longitude"] = lon if lon is not None else ""
                record["scraped_at"] = datetime.now(timezone.utc).isoformat()

            append_jsonl(record)
            scraped_this_run += 1
            elapsed = time.time() - t0

            remaining_after = remaining - i
            status = record["scrape_status"]
            name = record.get("factory_name") or "(no name)"
            print(f"[{i}/{remaining}] id={fid} status={status} ({elapsed:.1f}s) - {name} "
                  f"| remaining: {remaining_after}")

            if scraped_this_run % 20 == 0:
                try:
                    build_excel_from_jsonl()
                    print(f"    -> checkpoint saved to {OUTPUT_XLSX.name}")
                except PermissionError:
                    print(f"    -> checkpoint skipped: {OUTPUT_XLSX.name} is open elsewhere "
                          "(e.g. in Excel) -- please close it. Raw progress is still saved.")

            time.sleep(POLITE_DELAY_SECONDS)

    except KeyboardInterrupt:
        print("\nStopped by user. Progress is saved -- re-run this script to resume.")
    finally:
        print("\nBuilding final Excel file...")
        for attempt in range(1, 6):
            try:
                build_excel_from_jsonl()
                print(f"Saved {OUTPUT_XLSX}")
                break
            except PermissionError:
                print(f"{OUTPUT_XLSX.name} is open elsewhere (e.g. in Excel) -- "
                      f"please close it. Retrying in 10s... (attempt {attempt}/5)")
                time.sleep(10)
        else:
            print(f"Could not save {OUTPUT_XLSX.name} after 5 attempts -- close the file "
                  "and re-run this script; all scraped data is safe in "
                  f"{JSONL_PATH.relative_to(ROOT)}.")
        print(f"This run: {scraped_this_run} scraped, {login_pauses} login pause(s).")


if __name__ == "__main__":
    main()
